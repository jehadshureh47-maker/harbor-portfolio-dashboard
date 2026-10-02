#!/usr/bin/env python3
"""
verify_awqaf_month.py - Cross-check a month's Awqaf dashboard numbers against
the raw source workbooks.

WHY THIS EXISTS
----------------
On 2026-10-02 the September Awqaf upload shipped with two real bugs that
hand-checking a screenshot didn't catch:
  1. New Lease's INITIAL_SUMMARY (the Summary-sheet snapshot used to compute
     "Total New Leases" / "Total Rent Value") was accidentally left over from
     August's build, so the September numbers silently rendered as 0.
  2. Bounced Cheques' Unit Type column was built as all `null` instead of the
     real Residential/Commercial values, so that donut showed 100% Unknown.

Neither bug broke a SUM (the fee/amount totals were all correct), which is
why they were easy to miss just eyeballing totals. This script independently
recomputes every number from the raw .xlsx files and diffs it against what's
actually baked into Dashboard_v2.html for that month, so these two bug
classes - stale cross-month leftovers, and silently-null columns - get
caught automatically instead of by a user screenshot.

USAGE
-----
    python3 scripts/verify_awqaf_month.py <month-code> <raw-folder> [dashboard.html]

    month-code   e.g. 2026-10  (must be a key in the dashboard's
                 MONTH_DASHBOARD_SETS map)
    raw-folder   folder containing that month's 4 Awqaf workbooks
                 (Bounce Cheque / New Lease / Cheque Count / Request Count -
                 matched by keyword in the filename, case-insensitive)
    dashboard.html  defaults to Dashboard_v2.html next to this script's repo root

Exit code is 0 if every check passes, 1 if anything is flagged.
"""
import sys, re, json, base64, glob, os
from collections import Counter, defaultdict
from datetime import datetime

TOLERANCE = 0.01  # AED

# ---------------------------------------------------------------- utilities

def find_file(folder, *keywords):
    candidates = []
    for path in glob.glob(os.path.join(folder, "*.xlsx")):
        name = os.path.basename(path).lower()
        if all(kw.lower() in name for kw in keywords):
            candidates.append(path)
    if not candidates:
        raise FileNotFoundError(f"No .xlsx in {folder} matching keywords {keywords}")
    if len(candidates) > 1:
        print(f"  WARNING: multiple files match {keywords}, using first: {candidates}")
    return candidates[0]


def load_sheet(path, sheet=None):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    header = rows[0]
    data = [r for r in rows[1:] if any(c is not None for c in r)]
    return header, data


def get_month_object_name(dashboard_html, month_code):
    m = re.search(r"MONTH_DASHBOARD_SETS\s*=\s*\{(.*?)\};", dashboard_html, re.DOTALL)
    if not m:
        raise RuntimeError("Could not find MONTH_DASHBOARD_SETS in dashboard HTML")
    block = m.group(1)
    entry = re.search(re.escape(f"'{month_code}'") + r"\s*:\s*\(typeof (\w+)", block)
    if not entry:
        raise RuntimeError(
            f"Month '{month_code}' is not registered in MONTH_DASHBOARD_SETS yet. "
            f"Add it there (and to MONTH_LABELS / MONTH_LABELS_AR / portfolioHasMonth) "
            f"before running this check."
        )
    return entry.group(1)  # e.g. EMBEDDED_DASHBOARDS_OCT


def extract_embedded_html(dashboard_html, object_name, key):
    idx = dashboard_html.find(f"const {object_name}")
    if idx == -1:
        raise RuntimeError(f"{object_name} not found in dashboard HTML")
    m = re.search(re.escape(f'"{key}": "'), dashboard_html[idx:idx + 200_000_000])
    if not m:
        raise RuntimeError(f'"{key}" not found inside {object_name}')
    start = idx + m.end()
    end = dashboard_html.find('"', start)
    b64 = dashboard_html[start:end]
    return base64.b64decode(b64).decode("utf-8")


def extract_js_literal(html, var_name, open_char, close_char):
    """Extract `const <var_name> = <literal>;` where literal starts with
    open_char/close_char (balanced bracket depth)."""
    m = re.search(re.escape(f"const {var_name} = "), html)
    if not m:
        return None
    start = m.end()
    depth = 0
    i = start
    while i < len(html):
        if html[i] == open_char:
            depth += 1
        elif html[i] == close_char:
            depth -= 1
            if depth == 0:
                return html[start:i + 1]
        i += 1
    return None


def js_object_to_json(lit):
    """Some templates write INITIAL_COLMAP with bare identifier keys
    (JS shorthand: {date:0, contract:1}) instead of quoted JSON keys
    ({"date":0}). Normalize bare keys to quoted ones before json.loads."""
    return re.sub(r'([{,]\s*)([A-Za-z_$][A-Za-z0-9_$]*)(\s*:)', r'\1"\2"\3', lit)


def parse_report_template(html):
    rows_lit = extract_js_literal(html, "INITIAL_ROWS", "[", "]")
    colmap_lit = extract_js_literal(html, "INITIAL_COLMAP", "{", "}")
    summary_lit = extract_js_literal(html, "INITIAL_SUMMARY", "{", "}")
    rows = json.loads(rows_lit) if rows_lit else None
    colmap = json.loads(js_object_to_json(colmap_lit)) if colmap_lit else None
    summary = json.loads(summary_lit) if summary_lit else None
    return rows, colmap, summary


# ---------------------------------------------------------- null-column / staleness checks

def check_null_columns(report_name, rows, colmap, issues):
    """Flag any mapped column that is 100% null/blank across all data rows.
    Catches the Bounced-Cheques-unitType-was-all-null class of bug."""
    if not rows or len(rows) < 2:
        return
    header = rows[0]
    data = rows[1:]
    for field, idx in colmap.items():
        if idx is None or idx < 0 or idx >= len(header):
            continue
        values = [r[idx] for r in data]
        non_null = [v for v in values if v not in (None, "", "null")]
        if len(data) > 0 and len(non_null) == 0:
            issues.append(
                f"[{report_name}] CRITICAL: column '{field}' ({header[idx]!r}) is blank/null "
                f"in ALL {len(data)} embedded rows - this exact pattern caused the September "
                f"Bounced Cheques Unit Type bug. Check whether this data was dropped when the "
                f"month's rows were built."
            )


def check_stale_summary(report_name, details_contracts, summary, issues):
    """Flag an INITIAL_SUMMARY whose contract keys don't overlap with this
    month's Details sheet contracts - the New Lease September bug."""
    if not summary or not summary.get("byContract"):
        return
    summary_contracts = set(summary["byContract"].keys())
    overlap = details_contracts & summary_contracts
    if summary_contracts and not overlap:
        issues.append(
            f"[{report_name}] CRITICAL: INITIAL_SUMMARY's {len(summary_contracts)} contract(s) "
            f"share ZERO overlap with this month's {len(details_contracts)} Details contracts. "
            f"This is almost certainly a leftover summary from a previous month (the exact bug "
            f"that zeroed out September's Total New Leases / Total Rent Value). Rebuild "
            f"INITIAL_SUMMARY from THIS month's own Summary sheet."
        )
    elif summary_contracts and len(overlap) < len(details_contracts):
        missing = details_contracts - summary_contracts
        issues.append(
            f"[{report_name}] WARNING: {len(missing)} of {len(details_contracts)} Details "
            f"contracts have no matching entry in INITIAL_SUMMARY: {sorted(missing)[:10]}"
            + (" ..." if len(missing) > 10 else "")
        )


def compare(report_name, metric, raw_val, embedded_val, issues, tol=TOLERANCE):
    if isinstance(raw_val, (int, float)) and isinstance(embedded_val, (int, float)):
        ok = abs(raw_val - embedded_val) <= tol
    else:
        ok = raw_val == embedded_val
    status = "OK " if ok else "FAIL"
    print(f"  [{status}] {metric}: raw={raw_val!r}  embedded={embedded_val!r}")
    if not ok:
        issues.append(f"[{report_name}] MISMATCH on '{metric}': raw workbook says {raw_val!r}, "
                       f"dashboard shows {embedded_val!r}")


# --------------------------------------------------------------- per-report checks

def verify_bounced_cheques(folder, dashboard_html, object_name, issues):
    print("\n=== Bounced Cheques ===")
    path = find_file(folder, "bounce")
    header, data = load_sheet(path, sheet="Details")
    amount_idx = header.index("Cheque Amount")
    ctype_idx = header.index("Contract Type")

    raw_total = round(sum(r[amount_idx] for r in data if r[amount_idx] is not None), 2)
    raw_count = len(data)
    raw_unit_split = Counter(r[ctype_idx] for r in data)

    html = extract_embedded_html(dashboard_html, object_name, "bounced-cheques")
    rows, colmap, _ = parse_report_template(html)
    check_null_columns("Bounced Cheques", rows, colmap, issues)

    edata = rows[1:]
    amt_i, ut_i = colmap["amount"], colmap["unitType"]
    emb_total = round(sum(r[amt_i] for r in edata if r[amt_i] is not None), 2)
    emb_count = len(edata)
    emb_unit_split = Counter((r[ut_i] or "Unknown") for r in edata)

    compare("Bounced Cheques", "row count", raw_count, emb_count, issues)
    compare("Bounced Cheques", "total amount", raw_total, emb_total, issues)
    compare("Bounced Cheques", "Residential count", raw_unit_split.get("Residential", 0),
            emb_unit_split.get("Residential", 0), issues)
    compare("Bounced Cheques", "Commercial count", raw_unit_split.get("Commercial", 0),
            emb_unit_split.get("Commercial", 0), issues)


def verify_new_lease(folder, dashboard_html, object_name, issues):
    print("\n=== New Lease ===")
    path = find_file(folder, "new", "lease")
    dheader, ddata = load_sheet(path, sheet="Details")
    sheader, sdata = load_sheet(path, sheet="Summary")

    amount_idx = dheader.index("Amount")
    fee_idx = dheader.index("Description")
    contract_idx = dheader.index("Contract Number")

    raw_total_fees = round(sum(r[amount_idx] for r in ddata if r[amount_idx] is not None), 2)
    details_contracts = set(str(r[contract_idx]).strip() for r in ddata if r[contract_idx])

    raw_summary = {}
    for r in sdata:
        if r[0] is None:
            continue
        raw_summary[str(r[0]).strip()] = r[4]  # Rent Amount column
    raw_rent_value_total = round(sum(v for v in raw_summary.values() if v is not None), 2)

    html = extract_embedded_html(dashboard_html, object_name, "new-lease")
    rows, colmap, summary = parse_report_template(html)
    check_stale_summary("New Lease", details_contracts, summary, issues)

    edata = rows[1:]
    amt_i = colmap["amount"]
    emb_total_fees = round(sum(r[amt_i] for r in edata if r[amt_i] is not None), 2)

    emb_contracts = set(summary["byContract"].keys()) if summary and summary.get("byContract") else set()
    linked = details_contracts & emb_contracts
    emb_new_leases = len(linked)
    emb_rent_value = round(sum(summary["byContract"][c]["rentAmount"] for c in linked), 2) if summary else 0

    compare("New Lease", "total fees recorded", raw_total_fees, emb_total_fees, issues)
    compare("New Lease", "distinct new-lease contracts (Details x Summary link)",
            len(details_contracts), emb_new_leases, issues)
    compare("New Lease", "total rent value (Summary sum)", raw_rent_value_total, emb_rent_value, issues)


def verify_cheque_count(folder, dashboard_html, object_name, issues):
    print("\n=== Cheque Count ===")
    path = find_file(folder, "cheque", "count")
    header, data = load_sheet(path, sheet="Details")

    amount_idx = header.index("Amount")
    leasetype_idx = header.index("Type of lease")
    contract_idx = header.index("Contract Number")

    raw_total = round(sum(r[amount_idx] for r in data if r[amount_idx] is not None), 2)

    contract_type = {}
    for r in data:
        c = str(r[contract_idx]).strip() if r[contract_idx] else None
        if c and c not in contract_type:
            contract_type[c] = r[leasetype_idx]
    raw_lease_counts = Counter(contract_type.values())

    html = extract_embedded_html(dashboard_html, object_name, "count-report")
    rows, colmap, _ = parse_report_template(html)
    check_null_columns("Cheque Count", rows, colmap, issues)

    edata = rows[1:]
    amt_i = colmap["amount"]
    emb_total = round(sum(r[amt_i] for r in edata if r[amt_i] is not None), 2)

    compare("Cheque Count", "row count", len(data), len(edata), issues)
    compare("Cheque Count", "total amount", raw_total, emb_total, issues)
    print(f"  (reference) raw distinct contracts (trimmed): {len(contract_type)}, "
          f"lease-type split: {dict(raw_lease_counts)}")


def verify_request_count(folder, dashboard_html, object_name, issues):
    print("\n=== Request Count ===")
    path = find_file(folder, "request", "count")
    header, data = load_sheet(path, sheet="Details")

    amount_idx = header.index("Amount")

    raw_total = round(sum(r[amount_idx] for r in data if r[amount_idx] is not None), 2)
    raw_count = len(data)

    html = extract_embedded_html(dashboard_html, object_name, "request-count")
    rows, colmap, _ = parse_report_template(html)
    check_null_columns("Request Count", rows, colmap, issues)

    edata = rows[1:]
    amt_i = colmap["amount"]
    emb_total = round(sum(r[amt_i] for r in edata if r[amt_i] is not None), 2)
    emb_count = len(edata)

    compare("Request Count", "row count", raw_count, emb_count, issues)
    compare("Request Count", "total value", raw_total, emb_total, issues)
    compare("Request Count", "average per request", round(raw_total / raw_count, 2) if raw_count else 0,
            round(emb_total / emb_count, 2) if emb_count else 0, issues)


# --------------------------------------------------------------------- main

def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    month_code = sys.argv[1]
    folder = sys.argv[2]
    repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dashboard_path = sys.argv[3] if len(sys.argv) > 3 else os.path.join(repo_root, "Dashboard_v2.html")

    with open(dashboard_path, encoding="utf-8") as f:
        dashboard_html = f.read()

    object_name = get_month_object_name(dashboard_html, month_code)
    print(f"Month '{month_code}' -> {object_name}")

    issues = []
    verify_bounced_cheques(folder, dashboard_html, object_name, issues)
    verify_new_lease(folder, dashboard_html, object_name, issues)
    verify_cheque_count(folder, dashboard_html, object_name, issues)
    verify_request_count(folder, dashboard_html, object_name, issues)

    print("\n" + "=" * 70)
    if issues:
        print(f"{len(issues)} ISSUE(S) FOUND:\n")
        for i in issues:
            print(" -", i)
        print("\nFAIL - do not consider this month's data verified until these are fixed.")
        sys.exit(1)
    else:
        print("ALL CHECKS PASSED - every recomputed total matches the embedded dashboard data,")
        print("no null-column or stale cross-month data detected.")
        sys.exit(0)


if __name__ == "__main__":
    main()
